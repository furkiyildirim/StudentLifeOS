import os
import shutil
import csv
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QListWidget,
    QListWidgetItem, QTextEdit, QPushButton, QFileDialog,
    QMessageBox, QTabWidget, QScrollArea, QStackedWidget, 
    QComboBox, QFrame, QDialog, QLineEdit, QColorDialog, 
    QFontComboBox, QTableWidget, QTableWidgetItem, QApplication
)
from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import (
    QPixmap, QCursor, QFont, QTextCharFormat, 
    QTextListFormat
)
from PySide6.QtWebEngineWidgets import QWebEngineView
from PySide6.QtWebEngineCore import QWebEngineSettings

from core.sound import play_action_sound
from core.events import bus

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import pandas as pd
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

VAULT_DIR = "vault_storage"

# =========================================================================
# ANA EKRANA GÖMÜLÜ EXCEL/CSV TABLO EDİTÖRÜ (GELİŞMİŞ SÜRÜM)
# =========================================================================
class SpreadsheetEditorWidget(QWidget):
    def __init__(self, db, vault_view):
        super().__init__()
        self.db = db
        self.vault_view = vault_view
        self.current_file_path = None
        self.current_material_id = None
        self.init_ui()

    def init_ui(self):
        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        
        # --- ARAÇ ÇUBUĞU ---
        top_lay = QHBoxLayout()
        self.file_name_input = QLineEdit()
        self.file_name_input.setPlaceholderText("Dosya Adı (Örn: giderler.csv)")
        self.file_name_input.setStyleSheet("background-color: #171412; padding: 8px; border-radius: 6px; color: white; border: 1px solid #3f3f46;")
        
        btn_copy = QPushButton("📄 Kopyala")
        btn_paste = QPushButton("📋 Yapıştır")
        btn_add_row = QPushButton("+ Satır")
        btn_add_col = QPushButton("+ Sütun")
        btn_del_row = QPushButton("- Satır")
        btn_del_col = QPushButton("- Sütun")
        btn_clear = QPushButton("🧹 Temizle")
        
        for btn in [btn_copy, btn_paste, btn_add_row, btn_add_col, btn_del_row, btn_del_col, btn_clear]:
            btn.setCursor(QCursor(Qt.PointingHandCursor))
            btn.setStyleSheet("background-color: #27272a; color: white; padding: 6px; border-radius: 6px; border: 1px solid #3f3f46;")
        
        btn_save = QPushButton("💾 Kaydet")
        btn_save.setCursor(QCursor(Qt.PointingHandCursor))
        btn_save.setStyleSheet("background-color: #059669; color: white; padding: 8px 16px; font-weight: bold; border-radius: 6px;")
        
        top_lay.addWidget(self.file_name_input)
        top_lay.addWidget(btn_copy)
        top_lay.addWidget(btn_paste)
        top_lay.addWidget(btn_add_row)
        top_lay.addWidget(btn_add_col)
        top_lay.addWidget(btn_del_row)
        top_lay.addWidget(btn_del_col)
        top_lay.addWidget(btn_clear)
        top_lay.addWidget(btn_save)
        lay.addLayout(top_lay)

        # --- TABLO IZGARASI ---
        self.table = QTableWidget(15, 6)
        self.table.setStyleSheet("""
            QTableWidget { background-color: #ffffff; color: #000000; gridline-color: #d4d4d8; font-size: 13px; }
            QTableWidget QLineEdit { color: #000000; background-color: #ffffff; selection-background-color: #93c5fd; }
            QHeaderView::section { background-color: #f4f4f5; color: #18181b; font-weight: bold; padding: 4px; border: 1px solid #d4d4d8; }
        """)
        lay.addWidget(self.table)

        # --- FONKSİYON BAĞLANTILARI ---
        def insert_row():
            r = self.table.currentRow()
            self.table.insertRow(r if r >= 0 else self.table.rowCount())

        def insert_col():
            c = self.table.currentColumn()
            self.table.insertColumn(c if c >= 0 else self.table.columnCount())
            self.refresh_headers()
        
        def remove_row():
            r = self.table.currentRow()
            if r >= 0: self.table.removeRow(r)
            
        def remove_col():
            c = self.table.currentColumn()
            if c >= 0: 
                self.table.removeColumn(c)
                self.refresh_headers()
            
        def clear_cells():
            for item in self.table.selectedItems():
                item.setText("")

        def copy_cells():
            selection = self.table.selectedIndexes()
            if not selection: return
            rows = sorted(index.row() for index in selection)
            columns = sorted(index.column() for index in selection)
            rowcount = rows[-1] - rows[0] + 1
            colcount = columns[-1] - columns[0] + 1
            table = [[''] * colcount for _ in range(rowcount)]
            for index in selection:
                row = index.row() - rows[0]
                column = index.column() - columns[0]
                table[row][column] = str(index.data() or "")
            stream = '\n'.join(['\t'.join(row) for row in table])
            QApplication.clipboard().setText(stream)

        def paste_cells():
            text = QApplication.clipboard().text()
            if not text: return
            rows = text.split('\n')
            current_row = self.table.currentRow()
            current_col = self.table.currentColumn()
            if current_row < 0: current_row = 0
            if current_col < 0: current_col = 0
            
            for r, row in enumerate(rows):
                cols = row.split('\t')
                for c, val in enumerate(cols):
                    if current_row + r < self.table.rowCount() and current_col + c < self.table.columnCount():
                        self.table.setItem(current_row + r, current_col + c, QTableWidgetItem(val))

        btn_add_row.clicked.connect(insert_row)
        btn_add_col.clicked.connect(insert_col)
        btn_del_row.clicked.connect(remove_row)
        btn_del_col.clicked.connect(remove_col)
        btn_clear.clicked.connect(clear_cells)
        btn_copy.clicked.connect(copy_cells)
        btn_paste.clicked.connect(paste_cells)
        btn_save.clicked.connect(self.save_file)

    def refresh_headers(self):
        self.table.setHorizontalHeaderLabels([chr(65 + c) if c < 26 else f"Col{c}" for c in range(self.table.columnCount())])

    def new_file(self):
        self.current_file_path = None
        self.current_material_id = None
        self.file_name_input.clear()
        self.file_name_input.setReadOnly(False)
        self.table.clear()
        self.table.setRowCount(15)
        self.table.setColumnCount(6)
        self.refresh_headers()

    def load_file(self, material_id, filepath, filename):
        self.current_file_path = filepath
        self.current_material_id = material_id
        self.file_name_input.setText(filename)
        self.file_name_input.setReadOnly(True) 
        
        try:
            data = []
            if filepath.lower().endswith('.csv'):
                with open(filepath, 'r', encoding='utf-8') as f:
                    reader = csv.reader(f)
                    data = list(reader)
            else:
                if EXCEL_AVAILABLE:
                    df = pd.read_excel(filepath, header=None, dtype=str)
                    df.fillna("", inplace=True)
                    data = df.values.tolist()
            
            if data:
                self.table.setRowCount(len(data))
                self.table.setColumnCount(max((len(row) for row in data), default=6))
                for r, row in enumerate(data):
                    for c, val in enumerate(row):
                        self.table.setItem(r, c, QTableWidgetItem(str(val)))
                        
            self.refresh_headers()
        except Exception as e:
            QMessageBox.warning(self, "Hata", f"Tablo okunamadı: {e}")

    def save_file(self):
        name = self.file_name_input.text().strip()
        if not name:
            QMessageBox.warning(self, "Uyarı", "Lütfen tablo için bir dosya adı girin.")
            return
            
        if not self.current_file_path:
            if not name.lower().endswith((".csv", ".xlsx")):
                name += ".csv"
            dest = os.path.join(VAULT_DIR, name)
            self.current_file_path = dest
        else:
            dest = self.current_file_path

        data = []
        for r in range(self.table.rowCount()):
            row_data = []
            for c in range(self.table.columnCount()):
                item = self.table.item(r, c)
                row_data.append(item.text() if item else "")
            data.append(row_data)

        try:
            if EXCEL_AVAILABLE and dest.lower().endswith('.xlsx'):
                df = pd.DataFrame(data)
                df.to_excel(dest, header=False, index=False)
            else:
                with open(dest, 'w', newline='', encoding='utf-8') as f:
                    writer = csv.writer(f)
                    writer.writerows(data)
            
            if not self.current_material_id:
                with self.db.get_connection() as conn:
                    cur = conn.cursor()
                    cur.execute("INSERT INTO materials (file_name, file_path, file_type) VALUES (?, ?, 'excel')", (name, dest))
                    self.current_material_id = cur.lastrowid
                    conn.commit()
                self.vault_view.load_materials()
            
            try: play_action_sound("save")
            except: pass
            QMessageBox.information(self, "Başarılı", "Tablo başarıyla kaydedildi.")
        except Exception as e:
            QMessageBox.critical(self, "Hata", f"Kaydedilemedi: {e}")


class VaultView(QWidget):
    def __init__(self, db):
        super().__init__()
        self.db = db
        if not os.path.exists(VAULT_DIR):
            os.makedirs(VAULT_DIR)
            
        self.active_note_id = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)

        tabs = QTabWidget()
        tabs.addTab(self.create_notes_tab(), "📝 Gelişmiş Not Defteri")
        tabs.addTab(self.create_materials_tab(), "📂 Yerel Dosya & Tablo Havuzu")
        layout.addWidget(tabs)

    # =========================================================================
    # 1. BÖLÜM: WORD BENZERİ ZENGİN METİN EDİTÖRÜ 
    # =========================================================================
    def create_notes_tab(self):
        tab = QWidget()
        lay = QHBoxLayout(tab)
        lay.setContentsMargins(0, 12, 0, 0)
        lay.setSpacing(14)

        left_box = QVBoxLayout()
        left_box.setSpacing(8)

        btn_new_note = QPushButton("+ Yeni Not")
        btn_new_note.setObjectName("AccentButton")
        btn_new_note.setCursor(QCursor(Qt.PointingHandCursor))
        btn_new_note.clicked.connect(self.new_note)

        btn_import_note = QPushButton("📥 İçe Aktar (TXT/MD/DOCX)")
        btn_import_note.setCursor(QCursor(Qt.PointingHandCursor))
        btn_import_note.setStyleSheet("background-color: #27272a; padding: 6px; border-radius: 4px; border: 1px solid #3f3f46;")
        btn_import_note.clicked.connect(self.import_note)

        self.notes_list = QListWidget()
        self.notes_list.setFixedWidth(250)
        self.notes_list.itemClicked.connect(self.open_note)

        btn_delete_note = QPushButton("🗑 Seçili Notu Sil")
        btn_delete_note.setCursor(QCursor(Qt.PointingHandCursor))
        btn_delete_note.setStyleSheet("""
            background-color: #450a0a; color: #f87171;
            border: 1px solid #7f1d1d; border-radius: 6px;
            padding: 6px 12px; font-weight: 600;
        """)
        btn_delete_note.clicked.connect(self.delete_active_note)

        left_box.addWidget(btn_new_note)
        left_box.addWidget(btn_import_note)
        left_box.addWidget(self.notes_list)
        left_box.addWidget(btn_delete_note)
        lay.addLayout(left_box)

        editor_box = QWidget()
        e_lay = QVBoxLayout(editor_box)
        e_lay.setContentsMargins(0, 0, 0, 0)
        e_lay.setSpacing(8)
        
        top_ctrl = QHBoxLayout()
        self.note_title_edit = QTextEdit()
        self.note_title_edit.setMaximumHeight(38)
        self.note_title_edit.setStyleSheet("font-size: 16px; font-weight: bold; background-color: #171412; border: 1px solid #292524;")
        self.note_title_edit.setPlaceholderText("Not Başlığı...")
        
        btn_export = QPushButton("📤 Dışa Aktar")
        btn_export.setCursor(QCursor(Qt.PointingHandCursor))
        btn_export.setStyleSheet("background-color: #0284c7; padding: 6px 12px; border-radius: 6px; font-weight: bold;")
        btn_export.clicked.connect(self.export_note)

        btn_save = QPushButton("💾 Kaydet")
        btn_save.setObjectName("AccentButton")
        btn_save.setCursor(QCursor(Qt.PointingHandCursor))
        btn_save.clicked.connect(self.save_note)
        
        top_ctrl.addWidget(self.note_title_edit)
        top_ctrl.addWidget(btn_export)
        top_ctrl.addWidget(btn_save)
        e_lay.addLayout(top_ctrl)

        toolbar_frame = QFrame()
        toolbar_frame.setStyleSheet("background-color: #1c1917; border-radius: 6px; padding: 4px;")
        t_lay = QVBoxLayout(toolbar_frame)
        t_lay.setContentsMargins(4, 4, 4, 4)
        t_lay.setSpacing(4)
        
        row1 = QHBoxLayout()
        self.font_cb = QFontComboBox()
        self.font_cb.setStyleSheet("background-color: #27272a; color: white;")
        self.font_cb.currentFontChanged.connect(self.change_font)
        
        self.size_cb = QComboBox()
        self.size_cb.addItems(["8", "10", "12", "14", "16", "18", "24", "36"])
        self.size_cb.setCurrentText("14")
        self.size_cb.setStyleSheet("background-color: #27272a; color: white;")
        self.size_cb.currentTextChanged.connect(self.change_font_size)
        
        btn_color = QPushButton("🎨 Renk")
        btn_color.setStyleSheet("padding: 4px 10px; background-color: #27272a; border-radius: 4px;")
        btn_color.clicked.connect(self.change_color)
        
        row1.addWidget(self.font_cb)
        row1.addWidget(self.size_cb)
        row1.addWidget(btn_color)
        row1.addStretch()
        t_lay.addLayout(row1)
        
        row2 = QHBoxLayout()
        btn_style = "padding: 4px 12px; background-color: #27272a; border-radius: 4px;"
        
        btn_bold = QPushButton("B")
        btn_bold.setStyleSheet(f"font-weight: bold; font-family: serif; {btn_style}")
        btn_bold.clicked.connect(self.toggle_bold)
        
        btn_italic = QPushButton("I")
        btn_italic.setStyleSheet(f"font-style: italic; font-family: serif; {btn_style}")
        btn_italic.clicked.connect(self.toggle_italic)
        
        btn_underline = QPushButton("U")
        btn_underline.setStyleSheet(f"text-decoration: underline; font-family: serif; {btn_style}")
        btn_underline.clicked.connect(self.toggle_underline)

        btn_left = QPushButton("⇤ Sola Hizala")
        btn_left.setStyleSheet(btn_style)
        btn_left.clicked.connect(lambda: self.editor.setAlignment(Qt.AlignLeft))
        
        btn_center = QPushButton("⇥ Orta")
        btn_center.setStyleSheet(btn_style)
        btn_center.clicked.connect(lambda: self.editor.setAlignment(Qt.AlignCenter))

        btn_bullet = QPushButton("• Liste")
        btn_bullet.setStyleSheet(btn_style)
        btn_bullet.clicked.connect(self.insert_bullet)

        row2.addWidget(btn_bold)
        row2.addWidget(btn_italic)
        row2.addWidget(btn_underline)
        row2.addSpacing(10)
        row2.addWidget(btn_left)
        row2.addWidget(btn_center)
        row2.addSpacing(10)
        row2.addWidget(btn_bullet)
        row2.addStretch()
        t_lay.addLayout(row2)

        e_lay.addWidget(toolbar_frame)

        self.editor = QTextEdit()
        self.editor.setPlaceholderText("Ders notlarınızı buraya yazın...")
        self.editor.setStyleSheet("""
            QTextEdit {
                background-color: #ffffff;
                color: #000000;
                border-radius: 8px;
                padding: 30px;
                font-size: 14px;
                line-height: 1.6;
            }
        """)
        e_lay.addWidget(self.editor)

        lay.addWidget(editor_box)
        self.load_notes()
        return tab

    def change_font(self, font):
        fmt = QTextCharFormat()
        fmt.setFontFamily(font.family())
        self.editor.mergeCurrentCharFormat(fmt)
        self.editor.setFocus()

    def change_font_size(self, size_str):
        fmt = QTextCharFormat()
        fmt.setFontPointSize(float(size_str))
        self.editor.mergeCurrentCharFormat(fmt)
        self.editor.setFocus()

    def toggle_bold(self):
        fmt = QTextCharFormat()
        fmt.setFontWeight(QFont.Bold if self.editor.fontWeight() != QFont.Bold else QFont.Normal)
        self.editor.mergeCurrentCharFormat(fmt)
        self.editor.setFocus()

    def toggle_italic(self):
        fmt = QTextCharFormat()
        fmt.setFontItalic(not self.editor.fontItalic())
        self.editor.mergeCurrentCharFormat(fmt)
        self.editor.setFocus()

    def toggle_underline(self):
        fmt = QTextCharFormat()
        fmt.setFontUnderline(not self.editor.fontUnderline())
        self.editor.mergeCurrentCharFormat(fmt)
        self.editor.setFocus()

    def change_color(self):
        color = QColorDialog.getColor()
        if color.isValid():
            fmt = QTextCharFormat()
            fmt.setForeground(color)
            self.editor.mergeCurrentCharFormat(fmt)
        self.editor.setFocus()

    def insert_bullet(self):
        cursor = self.editor.textCursor()
        cursor.insertList(QTextListFormat.ListDisc)
        self.editor.setFocus()

    def load_notes(self):
        self.notes_list.clear()
        with self.db.get_connection() as conn:
            cur = conn.cursor()
            cur.execute("SELECT id, title FROM notes ORDER BY updated_at DESC")
            for row in cur.fetchall():
                item = QListWidgetItem(row["title"])
                item.setData(Qt.UserRole, row["id"])
                self.notes_list.addItem(item)

    def open_note(self, item):
        note_id = item.data(Qt.UserRole)
        self.active_note_id = note_id
        with self.db.get_connection() as conn:
            cur = conn.cursor()
            cur.execute("SELECT title, content FROM notes WHERE id = ?", (note_id,))
            n = cur.fetchone()
            if n:
                self.note_title_edit.setText(n["title"])
                self.editor.setHtml(n["content"] or "")

    def new_note(self):
        self.active_note_id = None
        self.note_title_edit.clear()
        self.editor.clear()
        self.note_title_edit.setFocus()

    def save_note(self):
        title = self.note_title_edit.toPlainText().strip()
        if not title:
            QMessageBox.warning(self, "Uyarı", "Lütfen bir not başlığı belirleyin.")
            return
        
        content = self.editor.toHtml()

        with self.db.get_connection() as conn:
            cur = conn.cursor()
            if self.active_note_id:
                cur.execute("UPDATE notes SET title = ?, content = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                            (title, content, self.active_note_id))
            else:
                cur.execute("INSERT INTO notes (title, content) VALUES (?, ?)", (title, content))
                self.active_note_id = cur.lastrowid
            conn.commit()
            
        try: play_action_sound("save")
        except: pass
        bus.notes_changed.emit()
        self.load_notes()

    def delete_active_note(self):
        current_item = self.notes_list.currentItem()
        target_id = self.active_note_id or (current_item.data(Qt.UserRole) if current_item else None)
        if not target_id: return

        confirm = QMessageBox.question(self, "Sil", "Bu notu kalıcı olarak silmek istediğinize emin misiniz?", QMessageBox.Yes | QMessageBox.No)
        if confirm == QMessageBox.Yes:
            with self.db.get_connection() as conn:
                conn.cursor().execute("DELETE FROM notes WHERE id = ?", (target_id,))
                conn.commit()
            try: play_action_sound("delete")
            except: pass
            self.new_note()
            self.load_notes()
            bus.notes_changed.emit()

    def import_note(self):
        path, _ = QFileDialog.getOpenFileName(self, "Not İçe Aktar", "", "Metin ve Word Dosyaları (*.txt *.md *.docx)")
        if not path: return
        
        try:
            content = ""
            title = os.path.basename(path).rsplit('.', 1)[0]
            
            if path.lower().endswith(".docx"):
                if not DOCX_AVAILABLE:
                    QMessageBox.warning(self, "Hata", "Word dosyalarını okumak için 'python-docx' kütüphanesi gerekli.")
                    return
                doc = docx.Document(path)
                content = "<br>".join([p.text for p in doc.paragraphs if p.text.strip()])
            else:
                with open(path, 'r', encoding='utf-8') as f:
                    content = f.read()
                    
            with self.db.get_connection() as conn:
                conn.cursor().execute("INSERT INTO notes (title, content) VALUES (?, ?)", (title, content))
                conn.commit()
            
            self.load_notes()
            bus.notes_changed.emit()
            QMessageBox.information(self, "Başarılı", "Belge içe aktarıldı.")
        except Exception as e:
            QMessageBox.critical(self, "Hata", f"Dosya okunamadı: {e}")

    def export_note(self):
        title = self.note_title_edit.toPlainText().strip()
        
        if not title or not self.editor.toPlainText().strip():
            QMessageBox.warning(self, "Uyarı", "Dışa aktarılacak bir not bulunamadı.")
            return
            
        path, _ = QFileDialog.getSaveFileName(self, "Notu Dışa Aktar", f"{title}.docx", "Word Dosyası (*.docx);;Markdown (*.md);;Metin (*.txt)")
        if not path: return
        
        try:
            if path.lower().endswith(".docx"):
                if not DOCX_AVAILABLE:
                    QMessageBox.warning(self, "Hata", "Word formatında kaydetmek için 'python-docx' kütüphanesi gerekli.")
                    return
                doc = docx.Document()
                doc.add_paragraph(self.editor.toPlainText())
                doc.save(path)
            else:
                with open(path, 'w', encoding='utf-8') as f:
                    f.write(self.editor.toPlainText())
            QMessageBox.information(self, "Başarılı", "Not dışa aktarıldı.")
        except Exception as e:
            QMessageBox.critical(self, "Hata", f"Kaydetme başarısız: {e}")

    # =========================================================================
    # 2. BÖLÜM: YEREL MATERYAL HAVUZU (Yerleşik Tablo Editörü)
    # =========================================================================
    def create_materials_tab(self):
        tab = QWidget()
        lay = QHBoxLayout(tab)
        lay.setContentsMargins(0, 12, 0, 0)
        lay.setSpacing(16)

        left_box = QVBoxLayout()
        left_box.setSpacing(10)

        btn_new_spreadsheet = QPushButton("📊 Yeni Tablo Oluştur (CSV)")
        btn_new_spreadsheet.setCursor(QCursor(Qt.PointingHandCursor))
        btn_new_spreadsheet.setStyleSheet("background-color: #064e3b; color: #10b981; padding: 6px; border-radius: 4px; border: 1px solid #047857; font-weight: bold;")
        btn_new_spreadsheet.clicked.connect(self.open_spreadsheet_editor)

        btn_add = QPushButton("+ Dosya Yükle (PDF, PPTX, DOCX, Excel, Resim)")
        btn_add.setObjectName("AccentButton")
        btn_add.setCursor(QCursor(Qt.PointingHandCursor))
        btn_add.clicked.connect(self.dialog_upload_material)
        
        left_box.addWidget(btn_new_spreadsheet)
        left_box.addWidget(btn_add)

        filter_lay = QVBoxLayout()
        filter_lay.setSpacing(4)
        lbl_filter = QLabel("🔍 Ders Filtresi:")
        lbl_filter.setStyleSheet("color: #a1a1aa; font-size: 11px; font-weight: bold;")
        
        self.course_filter_cb = QComboBox()
        self.course_filter_cb.currentIndexChanged.connect(self.load_materials)
        
        filter_lay.addWidget(lbl_filter)
        filter_lay.addWidget(self.course_filter_cb)
        left_box.addLayout(filter_lay)

        self.materials_list = QListWidget()
        self.materials_list.setFixedWidth(280)
        self.materials_list.itemClicked.connect(self.preview_material)
        left_box.addWidget(self.materials_list)

        self.btn_toggle_preview = QPushButton("👁 Ön İzlemeyi Gizle")
        self.btn_toggle_preview.setCheckable(True)
        self.btn_toggle_preview.setCursor(QCursor(Qt.PointingHandCursor))
        self.btn_toggle_preview.setStyleSheet("background-color: #27272a; border: 1px solid #3f3f46; border-radius: 6px; padding: 6px; font-weight: bold;")
        self.btn_toggle_preview.clicked.connect(self.toggle_preview)
        left_box.addWidget(self.btn_toggle_preview)

        btn_export_mat = QPushButton("📤 Bilgisayara Kaydet (Dışa Aktar)")
        btn_export_mat.setCursor(QCursor(Qt.PointingHandCursor))
        btn_export_mat.setStyleSheet("background-color: #0284c7; color: white; border-radius: 6px; padding: 6px 12px; font-weight: bold;")
        btn_export_mat.clicked.connect(self.export_material)
        left_box.addWidget(btn_export_mat)

        btn_delete_material = QPushButton("🗑 Seçili Dosyayı Sil")
        btn_delete_material.setCursor(QCursor(Qt.PointingHandCursor))
        btn_delete_material.setStyleSheet("background-color: #450a0a; color: #f87171; border: 1px solid #7f1d1d; border-radius: 6px; padding: 6px 12px; font-weight: 600;")
        btn_delete_material.clicked.connect(self.delete_active_material)
        left_box.addWidget(btn_delete_material)

        lay.addLayout(left_box)

        # Viewer Stack: Görüntüleyici Ekranları Barındırır
        self.viewer_stack = QStackedWidget()
        self.viewer_stack.setStyleSheet("background-color: rgba(128, 128, 128, 0.1); border-radius: 8px;")

        self.empty_label = QLabel("Görüntülemek için soldaki listeden bir dosya seçin.")
        self.empty_label.setAlignment(Qt.AlignCenter)
        self.empty_label.setStyleSheet("color: #71717a; font-size: 14px;")
        self.viewer_stack.addWidget(self.empty_label)

        self.image_scroll = QScrollArea()
        self.image_scroll.setWidgetResizable(True)
        self.image_scroll.setAlignment(Qt.AlignCenter)
        self.image_label = QLabel()
        self.image_label.setAlignment(Qt.AlignCenter)
        self.image_scroll.setWidget(self.image_label)
        self.viewer_stack.addWidget(self.image_scroll)

        self.pdf_view = QWebEngineView()
        self.pdf_view.settings().setAttribute(QWebEngineSettings.PluginsEnabled, True)
        self.pdf_view.settings().setAttribute(QWebEngineSettings.PdfViewerEnabled, True)
        self.viewer_stack.addWidget(self.pdf_view)
        
        self.document_viewer = QWebEngineView()
        self.viewer_stack.addWidget(self.document_viewer)

        self.spreadsheet_editor = SpreadsheetEditorWidget(self.db, self)
        self.viewer_stack.addWidget(self.spreadsheet_editor)

        lay.addWidget(self.viewer_stack)

        self.refresh_course_filter()
        self.load_materials()
        return tab

    def open_spreadsheet_editor(self):
        self.materials_list.clearSelection()
        self.spreadsheet_editor.new_file()
        self.viewer_stack.setCurrentIndex(4)

    def toggle_preview(self):
        is_hidden = self.btn_toggle_preview.isChecked()
        self.viewer_stack.setHidden(is_hidden)
        self.btn_toggle_preview.setText("👁 Ön İzlemeyi Aç" if is_hidden else "👁 Ön İzlemeyi Gizle")

    def refresh_course_filter(self):
        self.course_filter_cb.blockSignals(True)
        self.course_filter_cb.clear()
        self.course_filter_cb.addItem("Tüm Dersler ve Genel Materyaller", None)
        
        with self.db.get_connection() as conn:
            cur = conn.cursor()
            cur.execute("SELECT id, code FROM courses")
            for c in cur.fetchall():
                self.course_filter_cb.addItem(c["code"], c["id"])
                
        self.course_filter_cb.blockSignals(False)

    def load_materials(self):
        self.materials_list.clear()
        selected_course_id = self.course_filter_cb.currentData()

        query = """
            SELECT m.id, m.file_name, m.file_path, m.file_type, c.code 
            FROM materials m
            LEFT JOIN courses c ON m.course_id = c.id
        """
        params = []
        if selected_course_id is not None:
            query += " WHERE m.course_id = ?"
            params.append(selected_course_id)
            
        query += " ORDER BY m.created_at DESC"

        with self.db.get_connection() as conn:
            cur = conn.cursor()
            cur.execute(query, params)
            for r in cur.fetchall():
                ext = r["file_type"]
                icon_tag = "📕" if ext == "pdf" else "📊" if ext == "pptx" else "📝" if ext == "docx" else "📗" if ext == "excel" else "🖼️"
                course_tag = f"[{r['code']}] " if r["code"] else "[Genel] "
                
                item = QListWidgetItem(f"{icon_tag} {course_tag}{r['file_name']}")
                item.setData(Qt.UserRole, dict(r))
                self.materials_list.addItem(item)

    def dialog_upload_material(self):
        dlg = QDialog(self)
        dlg.setWindowTitle("Yeni Materyal Yükle")
        dlg.resize(400, 180)
        lay = QVBoxLayout(dlg)

        course_box = QComboBox()
        course_box.addItem("Genel (Ders Bağımsız)", None)
        with self.db.get_connection() as conn:
            cur = conn.cursor()
            cur.execute("SELECT id, code, name FROM courses")
            for c in cur.fetchall():
                course_box.addItem(f"{c['code']} - {c['name']}", c['id'])

        file_lay = QHBoxLayout()
        file_path_in = QLineEdit()
        file_path_in.setReadOnly(True)
        file_path_in.setPlaceholderText("Henüz bir dosya seçilmedi...")
        
        btn_browse = QPushButton("Gözat...")
        file_lay.addWidget(file_path_in)
        file_lay.addWidget(btn_browse)

        selected_file = {"path": ""}

        def browse():
            f_path, _ = QFileDialog.getOpenFileName(
                dlg, "Materyal Seç", "",
                "Tüm Desteklenenler (*.pdf *.pptx *.ppt *.docx *.xlsx *.xls *.csv *.png *.jpg);;Excel Dosyaları (*.xlsx *.xls *.csv);;Word (*.docx);;Sunumlar (*.pptx *.ppt);;PDF (*.pdf)"
            )
            if f_path:
                selected_file["path"] = f_path
                file_path_in.setText(os.path.basename(f_path))

        btn_browse.clicked.connect(browse)

        btn_save = QPushButton("Sisteme Yükle")
        btn_save.setObjectName("AccentButton")

        lay.addWidget(QLabel("Bu dosya hangi derse ait?"))
        lay.addWidget(course_box)
        lay.addWidget(QLabel("Yüklenecek Dosya:"))
        lay.addLayout(file_lay)
        lay.addWidget(btn_save)

        def save():
            if not selected_file["path"]: return
            f_path = selected_file["path"]
            file_name = os.path.basename(f_path)
            dest_path = os.path.join(VAULT_DIR, file_name)

            try:
                if f_path != os.path.abspath(dest_path):
                    shutil.copyfile(f_path, dest_path)
            except Exception as e:
                QMessageBox.critical(dlg, "Hata", f"Kopyalanamadı: {e}")
                return

            ext = file_name.lower()
            if ext.endswith(".pdf"): file_type = "pdf"
            elif ext.endswith(".docx"): file_type = "docx"
            elif ext.endswith(".pptx") or ext.endswith(".ppt"): file_type = "pptx"
            elif ext.endswith((".xlsx", ".xls", ".csv")): file_type = "excel"
            else: file_type = "image"
            
            c_id = course_box.currentData()

            with self.db.get_connection() as conn:
                conn.cursor().execute("INSERT INTO materials (course_id, file_name, file_path, file_type) VALUES (?, ?, ?, ?)",
                            (c_id, file_name, dest_path, file_type))
                conn.commit()

            try: play_action_sound("save")
            except: pass
            dlg.accept()
            self.load_materials()

        btn_save.clicked.connect(save)
        dlg.exec()

    def export_material(self):
        item = self.materials_list.currentItem()
        if not item: return
        data = item.data(Qt.UserRole)
        f_path = data["file_path"]

        if not os.path.exists(f_path):
            QMessageBox.critical(self, "Hata", "Dosya yerel diskte bulunamadı.")
            return
            
        dest_path, _ = QFileDialog.getSaveFileName(self, "Dosyayı Kaydet", data["file_name"], "Tüm Dosyalar (*.*)")
        if dest_path:
            try:
                shutil.copy2(f_path, dest_path)
                QMessageBox.information(self, "Başarılı", "Dosya dışa aktarıldı.")
            except Exception as e:
                QMessageBox.critical(self, "Hata", f"Dışa aktarma başarısız oldu: {e}")

    def delete_active_material(self):
        item = self.materials_list.currentItem()
        if not item: return

        data = item.data(Qt.UserRole)
        confirm = QMessageBox.question(self, "Sil", f"'{data['file_name']}' kalıcı olarak silinecek. Emin misiniz?", QMessageBox.Yes | QMessageBox.No)

        if confirm == QMessageBox.Yes:
            with self.db.get_connection() as conn:
                conn.cursor().execute("DELETE FROM materials WHERE id = ?", (data["id"],))
                conn.commit()
            play_action_sound("delete")
            if os.path.exists(data["file_path"]):
                try: os.remove(data["file_path"])
                except: pass

            self.viewer_stack.setCurrentIndex(0)
            self.load_materials()

    def preview_material(self, item):
        data = item.data(Qt.UserRole)
        f_path = data["file_path"]
        f_type = data["file_type"]

        if not os.path.exists(f_path):
            self.empty_label.setText("Dosya diskte bulunamadı.")
            self.viewer_stack.setCurrentIndex(0)
            return

        if f_type == "image":
            pixmap = QPixmap(f_path)
            if not pixmap.isNull():
                self.image_label.setPixmap(pixmap.scaledToWidth(800, Qt.SmoothTransformation))
                self.viewer_stack.setCurrentIndex(1)
        elif f_type == "pdf":
            self.pdf_view.setUrl(QUrl.fromLocalFile(os.path.abspath(f_path)))
            self.viewer_stack.setCurrentIndex(2)
        elif f_type == "docx":
            if not DOCX_AVAILABLE:
                self.document_viewer.setHtml("<h3 style='color:white;'>DOCX görüntülemek için 'python-docx' kütüphanesini kurmalısınız.<br>Terminal: pip install python-docx</h3>")
            else:
                try:
                    doc = docx.Document(f_path)
                    html = f"""
                    <html><head><style>
                    body {{ background-color: #171412; color: #f4f4f5; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif; padding: 40px; line-height: 1.6; }}
                    .doc-card {{ background-color: #1c1917; border-radius: 12px; padding: 40px; border: 1px solid #292524; }}
                    h2 {{ color: #38bdf8; border-bottom: 1px solid #292524; padding-bottom: 10px; margin-top: 0; }}
                    </style></head><body>
                    <div class='doc-card'>
                    <h2>{data['file_name']}</h2>
                    """
                    for p in doc.paragraphs:
                        if p.text.strip():
                            html += f"<p>{p.text}</p>"
                    html += "</div></body></html>"
                    self.document_viewer.setHtml(html)
                except Exception as e:
                    self.document_viewer.setHtml(f"<h3 style='color:red;'>Dosya okunamadı: {e}</h3>")
            self.viewer_stack.setCurrentIndex(3)
        elif f_type == "pptx":
            if not PPTX_AVAILABLE:
                self.document_viewer.setHtml("<h3 style='color:white;'>PPTX görüntülemek için 'python-pptx' kütüphanesini kurmalısınız.<br>Terminal: pip install python-pptx</h3>")
            else:
                try:
                    prs = Presentation(f_path)
                    
                    # Modern JS ve CSS Slayt Gösterisi Motoru
                    html = f"""
                    <html><head><style>
                    body {{ margin: 0; background: #0c0a09; color: #f4f4f5; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif; display: flex; flex-direction: column; height: 100vh; overflow: hidden; }}
                    .header {{ padding: 12px 20px; background: #12100e; border-bottom: 1px solid #292524; font-weight: bold; color: #a1a1aa; font-size: 14px; text-align: center; }}
                    .slide-container {{ flex: 1; position: relative; display: flex; justify-content: center; align-items: center; padding: 20px; }}
                    
                    /* Slayt Kartı ve Geçiş Efekti (Fade-in / Scale) */
                    .slide {{ display: none; width: 100%; max-width: 900px; max-height: 100%; background: #1c1917; border: 1px solid #292524; border-radius: 12px; padding: 50px; box-sizing: border-box; overflow-y: auto; box-shadow: 0 10px 30px rgba(0,0,0,0.5); animation: slideAnim 0.4s cubic-bezier(0.25, 0.46, 0.45, 0.94); }}
                    .slide.active {{ display: block; }}
                    @keyframes slideAnim {{ from {{ opacity: 0; transform: scale(0.96) translateY(10px); }} to {{ opacity: 1; transform: scale(1) translateY(0); }} }}
                    
                    .slide-title {{ color: #38bdf8; font-size: 28px; font-weight: 800; margin-top: 0; margin-bottom: 24px; border-bottom: 2px solid #292524; padding-bottom: 16px; line-height: 1.3; }}
                    ul {{ padding-left: 28px; margin-top: 0; }}
                    li {{ margin-bottom: 12px; font-size: 18px; color: #e4e4e7; line-height: 1.5; }}
                    
                    /* Navigasyon Paneli */
                    .controls {{ height: 70px; background: #12100e; display: flex; justify-content: center; align-items: center; gap: 24px; border-top: 1px solid #292524; }}
                    button {{ background: #27272a; color: #fff; border: 1px solid #3f3f46; padding: 10px 24px; border-radius: 6px; font-weight: bold; cursor: pointer; font-size: 14px; transition: 0.2s; }}
                    button:hover {{ background: #38bdf8; color: #000; border-color: #38bdf8; }}
                    .indicator {{ color: #a1a1aa; font-weight: 800; font-size: 15px; min-width: 70px; text-align: center; }}
                    </style></head><body>
                    
                    <div class="header">📊 {data['file_name']}</div>
                    <div class="slide-container">
                    """
                    
                    # Slayt verilerini çıkartıp HTML içine yerleştirme
                    for i, slide in enumerate(prs.slides):
                        html += f"<div class='slide'>"
                        
                        title_shape = getattr(slide.shapes, 'title', None)
                        if title_shape and hasattr(title_shape, 'text') and title_shape.text.strip():
                            html += f"<div class='slide-title'>{title_shape.text.strip()}</div>"
                            
                        html += "<ul>"
                        for shape in slide.shapes:
                            if title_shape and shape == title_shape:
                                continue
                            if hasattr(shape, "text") and shape.text.strip():
                                for paragraph in shape.text_frame.paragraphs:
                                    text = paragraph.text.strip()
                                    if text:
                                        html += f"<li>{text}</li>"
                        html += "</ul></div>"
                        
                    html += """
                    </div>
                    <div class="controls">
                        <button onclick="changeSlide(-1)">◀ Önceki Slayt</button>
                        <span class="indicator" id="slideNum">1 / X</span>
                        <button onclick="changeSlide(1)">Sonraki Slayt ▶</button>
                    </div>
                    
                    <!-- Slayt Motoru JavaScript Kodları -->
                    <script>
                        let current = 0;
                        const slides = document.querySelectorAll('.slide');
                        const indicator = document.getElementById('slideNum');
                        
                        function showSlide(index) {
                            if (slides.length === 0) return;
                            if (index >= slides.length) current = 0;
                            if (index < 0) current = slides.length - 1;
                            
                            slides.forEach(s => s.classList.remove('active'));
                            slides[current].classList.add('active');
                            indicator.innerText = (current + 1) + ' / ' + slides.length;
                        }
                        
                        function changeSlide(dir) {
                            current += dir;
                            showSlide(current);
                        }
                        
                        // Klavye yön tuşları ile geçiş desteği
                        document.addEventListener('keydown', function(event) {
                            if (event.key === 'ArrowRight' || event.key === 'Space') { changeSlide(1); }
                            if (event.key === 'ArrowLeft') { changeSlide(-1); }
                        });
                        
                        showSlide(current);
                    </script>
                    </body></html>
                    """
                    self.document_viewer.setHtml(html)
                except Exception as e:
                    self.document_viewer.setHtml(f"<h3 style='color:red;'>Sunum okunamadı: {e}</h3>")
            self.viewer_stack.setCurrentIndex(3)
        elif f_type == "excel":
            self.spreadsheet_editor.load_file(data["id"], f_path, data["file_name"])
            self.viewer_stack.setCurrentIndex(4)